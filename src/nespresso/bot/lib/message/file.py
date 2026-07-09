import asyncio
import json
import os
import uuid
from typing import Any

import openpyxl
from aiogram import types

from nespresso.bot.lib.message.io import SendDocument
from nespresso.core.configs.paths import DIR_TEMP

# Cell values starting with these characters are interpreted as formulas by
# Excel/LibreOffice/Google Sheets when the file is opened — user- and
# upstream-controlled text (bios, hobbies, work history, …) flows straight into
# these exports, so any of them could otherwise plant a formula that executes
# on whoever opens the download (CSV/Excel formula injection).
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@")


def _SanitizeCell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + value
    return value


def ToJSONText(structure: dict[Any, Any] | list[dict[Any, Any]]) -> str:
    return json.dumps(structure, indent=3, ensure_ascii=False, default=str)


async def SendTemporaryFileFromText(chat_id: int, text: str) -> None:
    # uuid-suffixed so two concurrent exports for the same chat_id (a double-tap,
    # or overlapping admin actions) never share a path — otherwise one's
    # `finally: os.remove` would delete the file the other is still sending.
    file_path = DIR_TEMP / f"chat_id_{chat_id}_{uuid.uuid4().hex}.txt"

    with open(file_path, "w", encoding="utf-8") as file:
        file.write(text)

    try:
        await SendDocument(chat_id=chat_id, document=types.FSInputFile(file_path))
    finally:
        if file_path.exists():
            os.remove(file_path)


async def SendTemporaryXlsxFile(
    chat_id: int,
    sheets: list[tuple[str, list[str], list[list[Any]]]],
    filename: str = "export",
) -> None:
    """Build a multi-sheet xlsx workbook and send it as a document, then delete it."""
    # uuid-suffixed so two concurrent exports (a double-tap, or two admins
    # exporting at once) never share a path and race each other's write/delete.
    file_path = DIR_TEMP / f"{filename}_{chat_id}_{uuid.uuid4().hex}.xlsx"

    def _Build() -> None:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # type: ignore[arg-type]

        for sheet_name, headers, rows in sheets:
            ws = wb.create_sheet(title=sheet_name)
            ws.append(headers)
            for row in rows:
                ws.append([_SanitizeCell(v) for v in row])

        wb.save(file_path)

    # openpyxl is synchronous CPU/IO work; a large `message` table export would
    # otherwise block the single asyncio event loop that serves every user of
    # the bot for the whole duration of the export.
    await asyncio.to_thread(_Build)

    try:
        await SendDocument(chat_id=chat_id, document=types.FSInputFile(file_path))
    finally:
        if file_path.exists():
            os.remove(file_path)
